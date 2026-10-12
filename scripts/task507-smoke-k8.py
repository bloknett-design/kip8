#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Task 507: SMOKE kip8 (порт 9007, ключи localStorage БЕЗ префикса
# kip8test:) — перенос партии 504+505+506 одним инкрементом
# kipia-v525→v526:
#   A) Task 504: окно «Мероприятия» табеля — даты строк ЦВЕТНЫЕ
#      (выполнено → зелёная #81c784, просрочено → красная #ef5350,
#      фолбэк по дате, приоритет выполнения, гвард И/ПЗ);
#   B) Task 505/506: «Кабельный журнал» — кнопка печати, диалог,
#      страницы A4 с шапкой на каждой, сортировка № п/п, серые
#      «По проекту», Arial Narrow, Excel (объединения/Print_Titles);
#   C) 0 JS-ошибок; скриншоты.
import calendar
import datetime
import json
import os
import re
import threading
from http.server import HTTPServer, SimpleHTTPRequestHandler
from urllib.parse import unquote

from playwright.sync_api import sync_playwright

PORT = 9007
REPO = '/home/z/my-project/kip8'
SHOT_DIR = '/home/z/my-project/download/kip8-task507'
os.makedirs(SHOT_DIR, exist_ok=True)

PASS = 0
FAIL = 0
TODAY = datetime.date.today()
TODAY_ISO = TODAY.strftime('%Y-%m-%d')
Y, M = TODAY.year, TODAY.month
DIM = calendar.monthrange(Y, M)[1]


def d(off):
    dd = max(1, min(DIM, TODAY.day + off))
    return '%04d-%02d-%02d' % (Y, M, dd)


def check(name, cond, extra=''):
    global PASS, FAIL
    ok = bool(cond)
    if ok:
        PASS += 1
    else:
        FAIL += 1
    print(('  + ' if ok else '  X ') + name +
          (('  [' + str(extra)[:250] + ']') if (extra and not ok) else ''))


def shot(page, name):
    try:
        page.screenshot(path=os.path.join(SHOT_DIR, name))
    except Exception as e:
        print('  (скриншот %s: %s)' % (name, e))


# ---- мок А: Табель (Task 504) ----
CODES = [
  {'code': 'Д', 'name': 'День (12-час)', 'color': '#FFE082'},
  {'code': 'Н', 'name': 'Ночь (12-час)', 'color': '#B0BEC5'},
  {'code': 'И', 'name': 'Инструктаж', 'color': '#B3E5FC'},
  {'code': 'ОБ', 'name': 'Обучение', 'color': '#D1C4E9'},
  {'code': 'ПЗ', 'name': 'Проверка знаний', 'color': '#FFCDD2'},
]
EMPLOYEES = [
  {'таб_номер': '017', 'ФИО': 'Иванов Иван Иванович', 'тип': 'сменный',
   'смена': 1, 'шаблон_ротации': 1, 'старт_цикла': '%04d-%02d-01' % (Y, M),
   'дата_приёма': '2024-03-15', 'дата_увольнения': '', 'в_архиве': 0,
   'должность': 'Слесарь КИПиА', 'комментарий': ''},
  {'таб_номер': '023', 'ФИО': 'Петров Пётр Петрович', 'тип': 'дневной',
   'смена': '', 'шаблон_ротации': 2, 'старт_цикла': '%04d-%02d-07' % (Y, M),
   'дата_приёма': '2025-01-20', 'дата_увольнения': '', 'в_архиве': 0,
   'должность': 'Слесарь КИПиА', 'комментарий': ''},
]
PATTERNS = [
  {'id': 1, 'name': 'Сменный сутки/двое', 'cycle': 4, 'description': '',
   'days': [{'day': 1, 'status': 'Д'}, {'day': 2, 'status': 'Н'},
            {'day': 3, 'status': ''}, {'day': 4, 'status': ''}]},
]
ENTRIES = [
  {'id': 1, 'дата': d(-6), 'таб_номер': '017', 'статус': 'Д',
   'источник': 'авто'},
]
TRAININGS = [
  {'id': 100, 'тема': 'Повторный инструктаж по охране труда',
   'тип': 'инструктаж', 'дата_начала': d(-5), 'дата_окончания': d(-5),
   'таб_номер': '017', 'подразделение': '', 'выполнение': 1,
   'просрочен': 0},
  {'id': 101, 'тема': 'Проверка знаний до 1000В',
   'тип': 'проверка_знаний', 'дата_начала': d(-3),
   'дата_окончания': d(-3), 'таб_номер': '017', 'подразделение': '',
   'выполнение': 0, 'просрочен': 1},
  {'id': 102, 'тема': 'Инструктаж по пожарной безопасности',
   'тип': 'инструктаж', 'дата_начала': d(-4), 'дата_окончания': d(-4),
   'таб_номер': '023', 'подразделение': ''},
  {'id': 105, 'тема': 'Проверка знаний по электробезопасности',
   'тип': 'проверка_знаний', 'дата_начала': d(-1),
   'дата_окончания': d(-1), 'таб_номер': '023', 'подразделение': '',
   'выполнение': 1, 'просрочен': 1},
]
PPE = [
  {'id': 200, 'таб_номер': '017', 'наименование': 'Перчатки диэлектрические',
   'дата_выдачи': d(-60), 'дата_окончания': d(10), 'срок': '',
   'количество': 1},
]

# ---- мок Б: Кабельный журнал (Task 505/506) ----
CJ_COLUMNS = [
    {'key': 'designation', 'label': 'Обозначение кабеля, провода',
     'type': 'text'},
    {'key': 'start', 'label': 'Начало трассы', 'type': 'text'},
    {'key': 'end', 'label': 'Конец трассы', 'type': 'text'},
    {'key': 'section', 'label': 'Участок трассы кабеля, провода',
     'type': 'text'},
    {'key': 'mark_project', 'label': 'Марка (по проекту)', 'type': 'text'},
    {'key': 'cores_project',
     'label': 'Число и сечение жил (по проекту)', 'type': 'text'},
    {'key': 'length_project', 'label': 'Длина, м (по проекту)',
     'type': 'number'},
    {'key': 'mark_actual', 'label': 'Марка (проложен)', 'type': 'text'},
    {'key': 'cores_actual',
     'label': 'Число и сечение жил (проложен)', 'type': 'text'},
    {'key': 'length_actual', 'label': 'Длина, м (проложен)',
     'type': 'number'},
]
CJ_ROWS = [
    {'_row': 2, 'num': 1, 'designation': '1/7', 'start': 'ЩК-4, ШУ-1',
     'end': 'ЭЩ №4', 'section': 'по коробу',
     'mark_project': 'КВБбШЭнг-LS', 'cores_project': '14х1',
     'length_project': 60, 'mark_actual': 'КВБбШЭнг-LS',
     'cores_actual': '19х1', 'length_actual': 90,
     'department': 'Водоподготовка', 'purpose': 'Контрольный'},
    {'_row': 5, 'num': 3, 'designation': '5/2', 'start': 'ЩСУ-2',
     'end': 'Н-1', 'section': 'по лотку', 'mark_project': 'КГ',
     'cores_project': '1.0', 'length_project': 12, 'mark_actual': 'КГ',
     'cores_actual': '3х2,5', 'length_actual': 12,
     'department': 'Отделение №2', 'purpose': 'Силовой'},
    {'_row': 3, 'num': 2, 'designation': '2/1', 'start': 'ЩК-4, ШУ-1',
     'end': 'КС-1, отм. 0,0 м., оси 1-2/А-Б', 'section': 'по коробу',
     'mark_project': 'КВБбШЭнг-LS', 'cores_project': '27х1',
     'length_project': 43, 'mark_actual': 'КВБбШЭнг-LS',
     'cores_actual': '27х1', 'length_actual': 43.5,
     'department': 'Водоподготовка', 'purpose': 'Контрольный'},
]
for i in range(4, 31):
    CJ_ROWS.append({
        '_row': i + 1, 'num': i, 'designation': '%d/%d' % (i, i % 9 + 1),
        'start': 'ЩК-4, ШУ-%d' % (i % 7 + 1),
        'end': 'КС-%d, отм. %d,0 м.' % (i % 4 + 1, i % 18),
        'section': ['по коробу', 'по лотку', 'фальшпол', 'труба'][i % 4],
        'mark_project': 'КВБбШЭнг-LS', 'cores_project': '%dх1' % (i % 20 + 2),
        'length_project': (i * 7) % 180 + 5,
        'mark_actual': 'КВБбШЭнг-LS', 'cores_actual': '%dх1' % (i % 20 + 2),
        'length_actual': (i * 7) % 180 + 5,
        'department': ['Водоподготовка', 'МДИ', 'КИП и А', 'Склад'][i % 4],
        'purpose': 'Контрольный'})


def mock_response(action, body):
    if action == 'getCurrentUser':
        return {'ok': True, 'data': {'userId': 1, 'email': 'u@t.local',
                                     'role': 'Админ'}}
    if action == 'getMyAccess':
        return {'ok': True, 'data': {'role': 'Админ', 'found': True,
                'permissions': {'workschedule.view': True,
                                'workschedule.edit': True}}}
    if action == 'heartbeat':
        return {'ok': True, 'data': {'ok': True}}
    if action == 'workSchedule.getStatusCodes':
        return {'ok': True, 'data': {'codes': CODES}}
    if action == 'workSchedule.listEmployees':
        return {'ok': True, 'data': {'employees': EMPLOYEES}}
    if action == 'workSchedule.getPatterns':
        return {'ok': True, 'data': {'patterns': PATTERNS}}
    if action == 'workSchedule.listTrainings':
        return {'ok': True, 'data': {'trainings': TRAININGS}}
    if action == 'workSchedule.listPpe':
        return {'ok': True, 'data': {'ppe': PPE}}
    if action == 'workSchedule.listEntries':
        return {'ok': True, 'data': {'entries': ENTRIES}}
    if action == 'workSchedule.listVacations':
        return {'ok': True, 'data': {'vacations': []}}
    if action == 'cableJournal.getColumns':
        return {'ok': True, 'data': {'columns': CJ_COLUMNS,
                                     'canEdit': True}}
    if action == 'cableJournal.list':
        return {'ok': True, 'data': {'rows': CJ_ROWS,
                                     'total': len(CJ_ROWS),
                                     'canEdit': True}}
    return {'ok': True, 'data': {'ok': True}}


class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


def setup_routes(ctx):
    def handle(route, request):
        action = ''
        if 'action=' in request.url:
            action = unquote(request.url.split('action=')[1].split('&')[0])
        body = {}
        if request.post_data:
            try:
                body = json.loads(request.post_data)
            except Exception:
                body = {}
        return route.fulfill(status=200,
            content_type='application/json; charset=utf-8',
            body=json.dumps(mock_response(action, body), ensure_ascii=False))
    ctx.route('**/exec?**', handle)
    ctx.route('**script.google.com/**', handle)
    ctx.route('**raw.githubusercontent.com/**', lambda r: r.fulfill(
        status=404, content_type='text/plain', body='nf'))
    ctx.route('**calendar.legalic.ru/**', lambda r: r.fulfill(
        status=404, content_type='text/plain', body='nf'))
    ctx.route('**isdayoff.ru**', lambda r: r.fulfill(
        status=404, content_type='text/plain', body='nf'))


def seed_ls(theme):
    # kip8: ключи БЕЗ префикса kip8test:
    return ("try{localStorage.removeItem('kip8_ws_cache_v1')}catch(e){};" +
            "try{localStorage.removeItem('kip8_my_access')}catch(e){};" +
            "try{localStorage.removeItem('kip8_session_token')}catch(e){};" +
            "localStorage.setItem('kip8_session_token','bc-t507');" +
            "localStorage.setItem('app-theme','" + theme + "');")


PANEL_STATE_JS = r"""(() => {
    const p = document.getElementById('wsEventsPanel');
    if (!p) return {missing: true};
    const rows = [...p.querySelectorAll('.ws-ep-item')];
    return {hidden: p.hidden,
            rows: rows.map(r => {
                const dt = r.querySelector('.ws-ep-date');
                return {
                    cls: dt ? dt.className.replace('ws-ep-date', '').trim() : '',
                    color: dt ? getComputedStyle(dt).color : '',
                    text: (r.querySelector('.ws-ep-text') || {}).textContent || ''
                };})};})"""

CJ_FRAME_JS = r"""(() => {
    const f = document.querySelector('#cjPrintPrevModal .wspprev-frame');
    if (!f) return {missing: true};
    try {
        const d = f.contentDocument;
        if (!d) return {noDoc: true};
        const pages = [...d.querySelectorAll('.cjpr-page')];
        const first = pages[0];
        const des = [];
        d.querySelectorAll('.cjpr-page tbody tr').forEach(tr => {
            des.push((tr.cells[0] || {}).textContent || '');
        });
        return {
            pages: pages.length,
            titleGone: !d.querySelector('.cjpr-title, .cjpr-head'),
            theadPerPage: pages.map(p =>
                p.querySelectorAll('thead tr').length),
            thRowspan: first ?
                first.querySelectorAll('th[rowspan="2"]').length : 0,
            projGray: first ?
                (getComputedStyle(first.querySelector('td.cjpr-c-proj'))
                 .backgroundColor || '') : '',
            desJoined: des.join(','),
            sub: ((document.querySelector(
                '#cjPrintPrevModal .wspprev-sub') || {}).textContent || '')
        };
    } catch (e) { return {err: String(e)}; }})"""


def find_row(state, marker):
    for r in state.get('rows', []):
        if marker in r.get('text', ''):
            return r
    return None


def main():
    os.chdir(REPO)
    server = HTTPServer(('127.0.0.1', PORT), QuietHandler)
    threading.Thread(target=server.serve_forever, daemon=True).start()

    with sync_playwright() as p:
        browser = p.chromium.launch()

        # ======== A: Task 504 — окно «Мероприятия» ========
        print('A) Task 504: окно «Мероприятия» — цветные даты')
        ctx = browser.new_context(viewport={'width': 1280, 'height': 720})
        page = ctx.new_page()
        errors = []
        page.on('pageerror', lambda e: errors.append(str(e)))
        page.on('dialog', lambda dg: dg.accept())
        ctx.add_init_script(seed_ls('dark'))
        setup_routes(ctx)
        page.goto('http://localhost:%d/index.html' % PORT)
        page.wait_for_timeout(2500)
        page.evaluate("navigateTo('work-schedule')")
        page.wait_for_timeout(3500)
        st = page.evaluate(PANEL_STATE_JS)
        check('окно мероприятий отрисовано',
              not st.get('missing') and not st.get('hidden'), st)
        r1 = find_row(st, 'Повторный инструктаж')
        check('И выполнено → ЗЕЛЁНАЯ (ws-ep-date-ok, rgb(129,199,132))',
              r1 and r1['cls'] == 'ws-ep-date-ok' and
              r1['color'] == 'rgb(129, 199, 132)', r1)
        r2 = find_row(st, 'Проверка знаний до 1000В')
        check('ПЗ просрочен → КРАСНАЯ (ws-ep-date-bad, rgb(239,83,80))',
              r2 and r2['cls'] == 'ws-ep-date-bad' and
              r2['color'] == 'rgb(239, 83, 80)', r2)
        r6 = find_row(st, 'электробезопасности')
        check('выполнение ПРИОРИТЕТНЕЕ просрочки → ЗЕЛЁНАЯ',
              r6 and r6['cls'] == 'ws-ep-date-ok' and
              r6['color'] == 'rgb(129, 199, 132)', r6)
        shot(page, 'a1-events-dates.png')
        ctx.close()

        # ======== B: Task 505/506 — печать журнала ========
        print('B) Task 505/506: печать «Кабельного журнала»')
        ctx2 = browser.new_context(viewport={'width': 1280, 'height': 720},
                                   accept_downloads=True)
        page2 = ctx2.new_page()
        errors2 = []
        page2.on('pageerror', lambda e: errors2.append(str(e)))
        ctx2.add_init_script(seed_ls('dark'))
        setup_routes(ctx2)
        page2.goto('http://localhost:%d/index.html' % PORT)
        page2.wait_for_timeout(2500)
        page2.evaluate("navigateTo('cable-journal-edit')")
        page2.wait_for_timeout(1800)
        btn = page2.evaluate(
            "() => { const b = document.getElementById('cjPrintBtn');"
            " return !!b && b.offsetParent !== null; }")
        check('кнопка печати в шапке (Task 505)', btn)
        page2.click('#cjPrintBtn')
        page2.wait_for_timeout(1200)
        fr = page2.evaluate(CJ_FRAME_JS)
        check('предпросмотр: НЕСКОЛЬКО страниц A4 (Task 506)',
              fr.get('pages', 0) >= 2, fr.get('pages'))
        check('шапка НА КАЖДОЙ странице (thead 2×N)',
              all(t == 2 for t in fr.get('theadPerPage', [])),
              fr.get('theadPerPage'))
        check('заголовка-шапки НЕТ («просто таблицу»)',
              fr.get('titleGone'))
        check('шапка шаблона: 4 одиночные колонки (rowspan)',
              fr.get('thRowspan') == 4, fr.get('thRowspan'))
        check('серые «По проекту» (rgb(242, 242, 242))',
              '242, 242, 242' in fr.get('projGray', ''),
              fr.get('projGray'))
        check('сортировка исходного файла: 1/7 → 2/1 → 5/2',
              fr.get('desJoined', '').startswith('1/7,2/1,5/2'),
              fr.get('desJoined', '')[:50])
        check('подзаголовок диалога «записей: 30» без «отделений»',
              'записей: 30' in fr.get('sub', '') and
              'отделений' not in fr.get('sub', ''), fr.get('sub'))
        shot(page2, 'b1-preview.png')

        # Excel: объединения + Print_Titles (openpyxl)
        try:
            with page2.expect_download(timeout=30000) as dl:
                page2.click('#cjPrintPrevModal .wspprev-xlsx')
            d = dl.value
            data = open(d.path(), 'rb').read()
            out = os.path.join(SHOT_DIR, 'task507.xlsx')
            with open(out, 'wb') as f:
                f.write(data)
            import openpyxl
            wb = openpyxl.load_workbook(out)
            ws = wb['Кабельный журнал']
            merges = sorted(str(m) for m in ws.merged_cells.ranges)
            need = sorted(['A1:A2', 'B1:B2', 'C1:C2', 'D1:D2', 'E1:G1',
                           'H1:J1'])
            check('Excel: объединения шапки шаблона', merges == need,
                  merges)
            check('Excel: Print_Titles (шапка на каждой странице)',
                  ws.print_title_rows in ('1:2', '$1:$2'),
                  ws.print_title_rows)
            check('Excel: закрепление A3', ws.freeze_panes == 'A3',
                  ws.freeze_panes)
            check('Excel: порядок № п/п (r3=1/7, r4=2/1, r5=5/2)',
                  ws['A3'].value == '1/7' and ws['A4'].value == '2/1' and
                  ws['A5'].value == '5/2',
                  (ws['A3'].value, ws['A4'].value, ws['A5'].value))
            check('Excel: Arial Narrow 12 + серые «По проекту»',
                  ws['A3'].font.name == 'Arial Narrow' and
                  (ws['E3'].fill.fgColor.rgb or '').endswith('F2F2F2'),
                  (ws['A3'].font.name, ws['E3'].fill.fgColor.rgb))
        except Exception as e:
            check('Excel скачан', False, str(e)[:200])
        page2.keyboard.press('Escape')
        page2.wait_for_timeout(400)
        st2 = page2.evaluate(
            "() => !document.getElementById('cjPrintPrevModal') &&"
            " !document.getElementById('cjPrintStyle')")
        check('Esc закрыл диалог + снял инжект-стиль', st2)
        check('0 JS-ошибок (табель + журнал)', len(errors) == 0 and
              len(errors2) == 0, (errors[:3], errors2[:3]))
        ctx2.close()
        browser.close()

    print('\nИТОГО: %d/%d (fail=%d)' % (PASS, PASS + FAIL, FAIL))
    return 1 if FAIL else 0


if __name__ == '__main__':
    import sys
    sys.exit(main())
